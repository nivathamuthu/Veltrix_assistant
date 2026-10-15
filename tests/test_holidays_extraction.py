"""Holiday extraction from parsed calendar tables (Phase 3c). Synthetic documents only
(plus one check against the real corpus, skipped when data/ is absent)."""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

import pytest

import app.ingestion.holidays as holidays_module
from app.db import (
    CALENDAR_NO_PUBLIC,
    CALENDAR_OK,
    CALENDAR_UNAVAILABLE,
    KIND_OPTIONAL,
    KIND_PUBLIC,
    KIND_UNKNOWN,
)
from app.ingestion.corpus import LoadedDocument
from app.ingestion.holidays import classify_kind, extract_holidays
from app.ingestion.manifest import ManifestEntry
from app.ingestion.parsers import ParsedDocument, Section, parse_document, render_table

HEADER = ["Date", "Day", "Holiday"]


def entry(doc_id="SYN-CAL", regions=("IN",), effective="2026-01-01", doc_type="calendar", access="all"):
    return ManifestEntry(
        doc_id=doc_id, file=f"{doc_id}.pdf", title="Synthetic calendar", doc_type=doc_type,
        regions=tuple(regions), effective_date=effective, access_level=access,
    )


def table(heading, rows, header=HEADER) -> Section:
    data = render_table([header, *rows] if header else rows, has_header=bool(header))
    return Section(heading, 1, data.text, "table", data.rows)


def text(heading, body) -> Section:
    return Section(heading, 1, body, "text")


def doc_of(e: ManifestEntry, *sections: Section) -> LoadedDocument:
    return LoadedDocument(e, ParsedDocument(path=Path(e.file), parser="pdf", sections=list(sections)))


def run(e, *sections):
    return extract_holidays([e], [doc_of(e, *sections)])


def by_kind(result, kind):
    return sorted({h.date for h in result.holidays if h.kind == kind})


PUBLIC_ROWS = [
    ["15 January 2026", "Thursday", "Alpha Day"],
    ["26 January 2026", "Monday", "Beta Day"],
]
OPTIONAL_ROWS = [["4 March 2026", "Wednesday", "Gamma Day"]]


# --- kind from the heading ---------------------------------------------------


def test_heading_public_vs_restricted_decides_the_kind():
    result = run(
        entry(),
        text("Calendar", "Intro."),
        table("1. Public holidays 2026", PUBLIC_ROWS),
        table("2. Restricted holidays 2026", OPTIONAL_ROWS),
    )
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-15", "2026-01-26"]
    assert by_kind(result, KIND_OPTIONAL) == ["2026-03-04"]
    assert by_kind(result, KIND_UNKNOWN) == []
    (calendar,) = result.calendars
    assert (calendar.status, calendar.public_count, calendar.optional_count) == (CALENDAR_OK, 2, 1)


@pytest.mark.parametrize(
    "heading, kind",
    [
        ("1. Public holidays 2026", KIND_PUBLIC),
        ("Paid holidays 2026", KIND_PUBLIC),
        ("Statutory bank holidays", KIND_PUBLIC),
        ("Public and company holidays 2026", KIND_PUBLIC),
        ("PUBLIC HOLIDAYS", KIND_PUBLIC),
        ("2. Restricted holidays 2026", KIND_OPTIONAL),
        ("Optional holidays", KIND_OPTIONAL),
        ("Floating holidays", KIND_OPTIONAL),
        ("Optional (non-public) holidays", KIND_OPTIONAL),
        ("Holidays 2026", KIND_UNKNOWN),  # says nothing
        ("Company closure days", KIND_UNKNOWN),
        ("Republic Day observances", KIND_UNKNOWN),  # "public" only inside another word
        ("Public and restricted holidays", KIND_UNKNOWN),  # contradictory
        ("Non-public days", KIND_UNKNOWN),  # a negation is not a public signal
        ("Unpaid holidays", KIND_UNKNOWN),
        (None, KIND_UNKNOWN),
        ("", KIND_UNKNOWN),
    ],
)
def test_kind_vocabulary(heading, kind):
    assert classify_kind(heading, None).kind == kind


def test_caption_decides_only_when_the_heading_is_silent():
    assert classify_kind("Holidays 2026", "You may pick two of the restricted days below.") == (
        holidays_module.KindDecision(KIND_OPTIONAL, "caption")
    )
    assert classify_kind("Holidays 2026", "The office closes on the public holidays below.").kind == KIND_PUBLIC
    # a decisive heading is never overruled by the caption
    assert classify_kind("Public holidays", "restricted days are listed elsewhere").kind == KIND_PUBLIC
    # a contradictory heading is not rescued by a caption either
    assert classify_kind("Public and optional holidays", "public holidays below") == (
        holidays_module.KindDecision(KIND_UNKNOWN, "conflict")
    )


def test_caption_is_the_text_directly_above_the_table_under_the_same_heading():
    rows = OPTIONAL_ROWS
    same = run(entry(), text("Holidays 2026", "Choose from the restricted holidays below."),
               table("Holidays 2026", rows))
    assert by_kind(same, KIND_OPTIONAL) == ["2026-03-04"]
    # text under a DIFFERENT heading is not this table's caption
    other = run(entry(), text("Elsewhere", "Choose from the restricted holidays below."),
                table("Holidays 2026", rows))
    assert by_kind(other, KIND_UNKNOWN) == ["2026-03-04"]
    # text that comes AFTER the table is not used
    after = run(entry(), table("Holidays 2026", rows), text("Holidays 2026", "restricted holidays"))
    assert by_kind(after, KIND_UNKNOWN) == ["2026-03-04"]


def test_undetermined_kind_is_logged_kept_as_unknown_and_never_public(caplog):
    with caplog.at_level(logging.WARNING):
        result = run(entry("SYN-UNK"), table("Holidays 2026", PUBLIC_ROWS))
    assert by_kind(result, KIND_PUBLIC) == []
    assert by_kind(result, KIND_UNKNOWN) == ["2026-01-15", "2026-01-26"]
    (calendar,) = result.calendars
    assert calendar.status == CALENDAR_NO_PUBLIC and calendar.public_count == 0
    assert "SYN-UNK" in caplog.text and "kind cannot be determined" in caplog.text
    assert "Alpha" not in caplog.text  # holiday names are content: never logged


def test_restricted_and_optional_rows_never_count_as_public():
    result = run(
        entry(),
        table("Public holidays", PUBLIC_ROWS),
        table("Restricted holidays", OPTIONAL_ROWS),
        table("Optional holidays", [["5 May 2026", "Tuesday", "Delta Day"]]),
        table("Misc", [["6 June 2026", "Saturday", "Echo Day"]]),
    )
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-15", "2026-01-26"]
    assert {h.date for h in result.holidays} - set(by_kind(result, KIND_PUBLIC)) == {
        "2026-03-04", "2026-05-05", "2026-06-06",
    }


# --- rows that cannot be read --------------------------------------------------


def test_a_row_with_an_unreadable_date_is_skipped_and_logged_by_doc_id_and_row_number(caplog):
    rows = [
        ["15 January 2026", "Thursday", "SecretNameOne"],
        ["sometime in spring", "-", "SecretNameTwo"],
        ["03/04/2026", "-", "SecretNameThree"],  # ambiguous: never guessed
        ["16 January", "Friday", "SecretNameFour"],  # no year
        ["26 January 2026", "Monday", "SecretNameFive"],
    ]
    with caplog.at_level(logging.WARNING):
        result = run(entry("SYN-BAD"), table("Public holidays", rows))
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-15", "2026-01-26"]
    assert result.calendars[0].skipped_rows == 3
    for row_no in (2, 3, 4):
        assert f"doc_id=SYN-BAD table=1 row={row_no}" in caplog.text
    assert "row=1:" not in caplog.text and "row=5:" not in caplog.text
    assert "SecretName" not in caplog.text and "spring" not in caplog.text


def test_a_table_whose_dates_are_mostly_unreadable_is_not_dropped_silently(caplog):
    """German month names are not understood: every row must show up as a skipped,
    logged row (by number), and the calendar must be flagged, never quietly empty."""
    rows = [
        ["15. Januar 2026", "Donnerstag", "SecretNameOne"],
        ["26. Januar 2026", "Montag", "SecretNameTwo"],
        ["3. April 2026", "Freitag", "SecretNameThree"],
    ]
    with caplog.at_level(logging.WARNING):
        result = run(entry("SYN-DE"), table("Public holidays", rows))
    assert result.holidays == []
    (calendar,) = result.calendars
    assert (calendar.status, calendar.skipped_rows) == (CALENDAR_NO_PUBLIC, 3)
    for row_no in (1, 2, 3):
        assert f"doc_id=SYN-DE table=1 row={row_no}" in caplog.text
    assert "SecretName" not in caplog.text


def test_a_headerless_table_that_says_it_lists_holidays_but_has_few_dates_warns(caplog):
    rows = [["Alpha Day", "15 January 2026"], ["Beta Day", "soon"], ["Gamma Day", "later"], ["Delta Day", "TBC"]]
    with caplog.at_level(logging.WARNING):
        result = run(entry("SYN-FEW"), table("Public holidays", rows, header=None))
    assert result.holidays == [] and not result.tables[0].holiday_table
    assert "SYN-FEW table=1" in caplog.text and "dates in only 1 of 4 rows" in caplog.text


def test_a_row_without_a_name_is_skipped():
    rows = [["15 January 2026", "Thursday", "Alpha Day"], ["26 January 2026", "Monday", ""]]
    result = run(entry(), table("Public holidays", rows))
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-15"]
    assert result.calendars[0].skipped_rows == 1


def test_document_info_box_is_never_mistaken_for_holidays(caplog):
    info_rows = [["Document ID", "SYN-CAL"], ["Version", "2026"], ["Effective date", "1 January 2026"],
                 ["Owner", "People Team"], ["Classification", "Internal"]]
    # a title heading says nothing about holidays: the box is quietly ignored
    with caplog.at_level(logging.WARNING):
        quiet = run(entry(), table("Calendar 2026", info_rows, header=None),
                    table("Public holidays", PUBLIC_ROWS))
    assert by_kind(quiet, KIND_PUBLIC) == ["2026-01-15", "2026-01-26"]
    assert all(h.name not in ("Effective date", "Version") for h in quiet.holidays)
    assert [t.holiday_table for t in quiet.tables] == [False, True]
    assert "dates in only" not in caplog.text  # no noise for a normal info box
    # even under a heading that says "public", one date in five rows is no holiday list
    loud = run(entry(), table("Public holiday calendar 2026", info_rows, header=None))
    assert loud.holidays == []


def test_a_stale_weekday_keeps_the_date_and_warns(caplog):
    rows = [["16 January 2026", "Thursday", "Alpha Day"]]  # 16 Jan 2026 is a Friday
    with caplog.at_level(logging.WARNING):
        result = run(entry("SYN-DAY"), table("Public holidays", rows))
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-16"]  # the date cell wins
    assert "doc_id=SYN-DAY table=1 row=1" in caplog.text and "weekday does not match" in caplog.text


def test_the_name_is_never_taken_from_the_day_column_or_guessed_between_columns():
    ambiguous = table("Public holidays", [["15 January 2026", "Alpha", "Beta"]], header=["Date", "Col A", "Col B"])
    result = run(entry(), ambiguous)
    assert result.holidays == [] and result.calendars[0].skipped_rows == 1
    # a header that says which column is the name resolves it
    labelled = table("Public holidays", [["15 January 2026", "Alpha", "Beta"]],
                     header=["Date", "Notes", "Holiday name"])
    assert [h.name for h in run(entry(), labelled).holidays] == ["Beta"]


def test_tables_without_a_header_row_work_by_position():
    result = run(entry(), table("Public holidays", [["15 January 2026", "Alpha Day"], ["26 January 2026", "Beta Day"]], header=None))
    assert by_kind(result, KIND_PUBLIC) == ["2026-01-15", "2026-01-26"]


def test_dates_in_other_columns_and_orders_are_found_by_content_not_position():
    result = run(entry(), table("Public holidays", [["Alpha Day", "Thu", "Jan 15, 2026"], ["Beta Day", "Mon", "2026-01-26"]],
                                header=["Holiday", "Day", "When"]))
    assert {h.date: h.name for h in result.holidays} == {"2026-01-15": "Alpha Day", "2026-01-26": "Beta Day"}


# --- metadata ---------------------------------------------------------------------


def test_rows_carry_the_manifest_region_effective_date_and_regions():
    e = entry("SYN-MULTI", regions=("IN", "US"), effective="2026-04-01")
    result = run(e, table("Public holidays", PUBLIC_ROWS))
    assert len(result.holidays) == 4  # two dates x two regions
    assert {h.region for h in result.holidays} == {"IN", "US"}
    for h in result.holidays:
        assert (h.doc_id, h.effective_date, h.regions, h.kind) == ("SYN-MULTI", "2026-04-01", ("IN", "US"), KIND_PUBLIC)
    assert result.calendars[0].public_count == 2  # distinct dates, not rows x regions


def test_a_date_listed_twice_keeps_one_row_and_the_stronger_kind(caplog):
    with caplog.at_level(logging.WARNING):
        result = run(
            entry("SYN-DUP"),
            table("Restricted holidays", [["15 January 2026", "Thursday", "Alpha Day"]]),
            table("Public holidays", [["15 January 2026", "Thursday", "Alpha Day"]]),
        )
    assert [(h.date, h.kind) for h in result.holidays] == [("2026-01-15", KIND_PUBLIC)]
    assert "SYN-DUP" in caplog.text and "listed more than once" in caplog.text


def test_a_holiday_on_a_weekend_is_still_stored_as_given():
    result = run(entry(), table("Public holidays", [["15 August 2026", "Saturday", "Weekend Day"]]))
    assert by_kind(result, KIND_PUBLIC) == ["2026-08-15"]


# --- calendars: availability and scope ----------------------------------------------


def test_only_calendar_documents_are_read():
    policy = entry("SYN-POL", doc_type="policy")
    result = extract_holidays([policy], [doc_of(policy, table("Public holidays", PUBLIC_ROWS))])
    assert result.calendars == [] and result.holidays == []


def test_a_calendar_that_did_not_parse_is_recorded_as_unavailable(caplog):
    e = entry("SYN-GONE")
    with caplog.at_level(logging.WARNING):
        result = extract_holidays([e], [])
    (calendar,) = result.calendars
    assert calendar.status == CALENDAR_UNAVAILABLE and calendar.public_count == 0
    assert "SYN-GONE" in caplog.text


def test_one_failing_calendar_does_not_stop_the_others(monkeypatch):
    good, bad = entry("SYN-GOOD"), entry("SYN-BAD")
    real = holidays_module._extract_calendar

    def flaky(e, sections):
        if e.doc_id == "SYN-BAD":
            raise RuntimeError("boom")
        return real(e, sections)

    monkeypatch.setattr(holidays_module, "_extract_calendar", flaky)
    sections = [table("Public holidays", PUBLIC_ROWS)]
    result = extract_holidays([bad, good], [doc_of(bad, *sections), doc_of(good, *sections)])
    assert {c.doc_id: c.status for c in result.calendars} == {
        "SYN-BAD": CALENDAR_UNAVAILABLE, "SYN-GOOD": CALENDAR_OK,
    }
    assert {h.doc_id for h in result.holidays} == {"SYN-GOOD"}


def test_a_calendar_with_no_holiday_table_has_no_public_holidays():
    result = run(entry(), text("Calendar", "Just prose, no table."))
    assert result.calendars[0].status == CALENDAR_NO_PUBLIC and result.holidays == []


def test_extraction_is_deterministic():
    sections = [table("Public holidays", PUBLIC_ROWS), table("Restricted holidays", OPTIONAL_ROWS)]
    assert run(entry(), *sections).holidays == run(entry(), *sections).holidays


# --- the same calendar in every real format ---------------------------------------------
# Different date styles per file: the parsers feed the extractor structured rows.

PUBLIC_EXPECTED = ["2026-01-15", "2026-01-26", "2026-08-15"]
OPTIONAL_EXPECTED = ["2026-03-04"]


def _write_html(path: Path) -> None:
    path.write_text(
        "<html><body><h1>Calendar</h1><p>Intro text for the calendar.</p>"
        "<h2>Public holidays 2026</h2><table>"
        "<tr><th>Date</th><th>Day</th><th>Holiday</th></tr>"
        "<tr><td>Thu, 15 Jan 2026</td><td>Thursday</td><td>Alpha Day</td></tr>"
        "<tr><td>2026-01-26</td><td>Monday</td><td>Beta Day</td></tr>"
        "<tr><td>August 15th, 2026</td><td>Saturday</td><td>Gamma Day</td></tr></table>"
        "<h2>Restricted holidays 2026</h2><table>"
        "<tr><th>Date</th><th>Day</th><th>Holiday</th></tr>"
        "<tr><td>4 March 2026</td><td>Wednesday</td><td>Delta Day</td></tr></table>"
        "</body></html>",
        encoding="utf-8",
    )


def _write_markdown(path: Path) -> None:
    path.write_text(
        "# Calendar\n\nIntro text for the calendar.\n\n## Public holidays 2026\n\n"
        "| Date | Day | Holiday |\n|---|---|---|\n"
        "| 15 January 2026 | Thursday | Alpha Day |\n"
        "| 26/01/2026 | Monday | Beta Day |\n"
        "| Saturday, 15 August 2026 | Saturday | Gamma Day |\n\n"
        "## Restricted holidays 2026\n\n| Date | Day | Holiday |\n|---|---|---|\n"
        "| 2026-03-04 | Wednesday | Delta Day |\n",
        encoding="utf-8",
    )


def _write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Calendar", level=1)
    document.add_paragraph("Intro text for the calendar.")

    def add_table(heading, rows):
        document.add_heading(heading, level=2)
        grid = document.add_table(rows=1 + len(rows), cols=3)
        for col, label in enumerate(("Date", "Day", "Holiday")):
            run_ = grid.rows[0].cells[col].paragraphs[0].add_run(label)
            run_.bold = True
        for r, values in enumerate(rows, start=1):
            for col, value in enumerate(values):
                grid.rows[r].cells[col].text = value

    add_table("Public holidays 2026", [
        ["15 Jan 2026", "Thursday", "Alpha Day"],
        ["Monday 26 January 2026", "Monday", "Beta Day"],
        ["15.08.2026", "Saturday", "Gamma Day"],
    ])
    add_table("Restricted holidays 2026", [["Mar 4, 2026", "Wednesday", "Delta Day"]])
    document.save(path)


def _write_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=8)
    heading = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=6)
    style = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 10)])

    def grid(rows):
        return Table([["Date", "Day", "Holiday"], *rows], style=style, hAlign="LEFT")

    story = [
        Paragraph("Calendar", title),
        Paragraph("Intro text for the calendar, long enough to be the body text size of this page.", body),
        Paragraph("Public holidays 2026", heading),
        grid([["15 January 2026", "Thursday", "Alpha Day"],
              ["Jan 26, 2026", "Monday", "Beta Day"],
              ["2026-08-15", "Saturday", "Gamma Day"]]),
        Spacer(1, 12),
        Paragraph("Restricted holidays 2026", heading),
        grid([["4 March 2026", "Wednesday", "Delta Day"]]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


@pytest.mark.parametrize(
    "suffix, writer",
    [(".html", _write_html), (".md", _write_markdown), (".docx", _write_docx), (".pdf", _write_pdf)],
)
def test_same_holidays_from_every_format_with_different_date_styles(tmp_path, suffix, writer):
    path = tmp_path / f"calendar{suffix}"
    writer(path)
    parsed = parse_document(path)
    assert parsed.ok
    e = entry("SYN-FMT")
    result = extract_holidays([e], [LoadedDocument(e, parsed)])
    assert by_kind(result, KIND_PUBLIC) == PUBLIC_EXPECTED
    assert by_kind(result, KIND_OPTIONAL) == OPTIONAL_EXPECTED
    assert by_kind(result, KIND_UNKNOWN) == []
    names = {h.date: h.name for h in result.holidays}
    assert names["2026-01-15"] == "Alpha Day" and names["2026-03-04"] == "Delta Day"


# --- the real corpus (skipped when data/ is absent) ---------------------------------------

DATA = Path(__file__).resolve().parent.parent / "data"


@pytest.mark.skipif(not (DATA / "manifest.json").is_file(), reason="data/ not present")
def test_real_calendars_extract_every_dated_row_and_only_public_rows_count():
    from app.ingestion.corpus import load_corpus

    corpus = load_corpus(DATA / "corpus", DATA / "manifest.json")
    result = extract_holidays(corpus.entries, corpus.documents)
    assert result.calendars, "the corpus has calendar documents"
    documents = {d.entry.doc_id: d.parsed for d in corpus.documents}
    for calendar in result.calendars:
        assert calendar.status == CALENDAR_OK and calendar.skipped_rows == 0
        # independent check against the text form: one 'Row: Date ...' line per holiday
        text_rows = sum(
            1
            for s in documents[calendar.doc_id].sections
            if s.kind == "table"
            for line in s.text.splitlines()
            if line.startswith("Row: Date ")
        )
        assert calendar.public_count + calendar.optional_count + calendar.unknown_count == text_rows
        assert calendar.unknown_count == 0
    # every optional row sits under a heading that says so, and is never public
    for h in result.holidays:
        assert h.kind in (KIND_PUBLIC, KIND_OPTIONAL)
