"""DOCX parser (python-docx).

Paragraphs and tables are read in document order. ``Title`` and ``Heading N``
paragraph styles are headings. Body paragraphs often have no style at all
(``paragraph.style`` can be ``None``), so that case is handled.
"""
from __future__ import annotations

from pathlib import Path
from typing import Iterator

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.parsers.builder import SectionBuilder, clean
from app.ingestion.parsers.models import ParsedDocument
from app.ingestion.parsers.tables import rows_to_text

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _is_heading(paragraph: Paragraph) -> bool:
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    name = name.lower()
    return name == "title" or name.startswith("heading")


def _blocks(parent_element, parent) -> Iterator[Paragraph | Table]:
    """Paragraphs and tables of ``parent_element`` in document order.

    Block-level content controls (``w:sdt``) are looked into, because they
    wrap ordinary paragraphs and tables.
    """
    for child in parent_element.iterchildren():
        if child.tag == _W + "p":
            yield Paragraph(child, parent)
        elif child.tag == _W + "tbl":
            yield Table(child, parent)
        elif child.tag == _W + "sdt":
            content = child.find(_W + "sdtContent")
            if content is not None:
                yield from _blocks(content, parent)


def _row_cells(row) -> list:
    """A row's distinct cells: horizontally merged cells repeat in python-docx."""
    cells, previous = [], None
    for cell in row.cells:
        if previous is not None and cell._tc is previous:
            continue
        cells.append(cell)
        previous = cell._tc
    return cells


def _is_bold(cell) -> bool:
    runs = [r for p in cell.paragraphs for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def _has_header(table: Table) -> bool:
    """First row is a header if Word marks it so, or if all its text is bold."""
    if not table.rows:
        return False
    first = table.rows[0]
    properties = first._tr.trPr
    if properties is not None and properties.find(_W + "tblHeader") is not None:
        return True
    cells = [c for c in _row_cells(first) if c.text.strip()]
    return bool(cells) and all(_is_bold(c) for c in cells)


def _table_text(table: Table) -> str:
    rows = [[cell.text for cell in _row_cells(row)] for row in table.rows]
    return rows_to_text(rows, _has_header(table))


def parse_docx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    document = Document(str(path))
    builder = SectionBuilder()
    for block in _blocks(document.element.body, document):
        if isinstance(block, Table):
            builder.table(_table_text(block))
            continue
        text = clean(block.text)
        if not text:
            continue
        if _is_heading(block):
            builder.heading(text)
        else:
            builder.paragraph(text)
    return ParsedDocument(path=path, parser="docx", sections=builder.finish())
